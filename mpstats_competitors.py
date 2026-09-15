# -*- coding: utf-8 -*-
"""
MPSTATS Ozon competitor collector
---------------------------------
Берет запросы из файла "запросы 62 SKU.xlsx", выбирает TOP-N запросов
на каждый наш SKU, получает выдачу через тот же endpoint MPSTATS, который
использует веб-интерфейс, и собирает Excel с конкурентами.

ВАЖНО:
- Логин/пароль и cookies в коде НЕ хранятся.
- При первом запуске откроется отдельный Chrome-профиль.
  Один раз войдите в MPSTATS вручную.
- Скрипт уважает 429 и Retry-After, не пытается обходить ограничения.
- Результаты запросов кэшируются в SQLite, поэтому после сбоя можно
  просто запустить скрипт снова.
"""

from __future__ import annotations

import json
import math
import re
import sqlite3
import threading
import time
from collections import defaultdict
from concurrent.futures import ThreadPoolExecutor, as_completed
from dataclasses import dataclass
from datetime import datetime, timezone
from pathlib import Path
from statistics import mean
from typing import Any
from urllib.parse import quote

import requests
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from selenium import webdriver
from selenium.webdriver.chrome.options import Options


# ============================================================
# НАСТРОЙКИ — обычно достаточно менять только этот блок
# ============================================================

BASE_DIR = Path(__file__).resolve().parent

INPUT_XLSX = BASE_DIR / "запросы 62 SKU.xlsx"
OUTPUT_XLSX = BASE_DIR / "mpstats_competitors.xlsx"
CACHE_DB = BASE_DIR / "mpstats_cache.sqlite3"
CHROME_PROFILE_DIR = BASE_DIR / "mpstats_profile"

# Быстрый режим: TOP-10 наиболее частотных запросов на каждый SKU.
# Если поставить 0 — брать все подходящие запросы.
TOP_QUERIES_PER_SKU = 10

# Берем строки, где "Тематический кандидат" начинается с "Да".
# Если для конкретного SKU после фильтра ничего не осталось,
# скрипт автоматически откатится к его полному списку запросов.
THEMATIC_ONLY = True

# Если True — дополнительно оставляем только строки:
# "Сопоставим: ≥7 дней в каждом" = "Да".
COMPARABLE_ONLY = False

# Сколько items сохранять из ответа MPSTATS.
# Endpoint обычно отдает до ~200. 0 = сохранить все.
MAX_ITEMS_PER_QUERY = 0

# Сколько конкурентов оставить в агрегированном листе на каждый наш SKU.
TOP_COMPETITORS_PER_SKU = 50

# Количество одновременных запросов. Начните с 4.
# Не повышайте агрессивно: при 429 скрипт сам замедлится.
MAX_WORKERS = 4

# Пауза перед запросом каждого worker, сек.
BASE_REQUEST_DELAY = 0.20

# Сколько раз повторять запрос при 429 / временной ошибке.
MAX_RETRIES = 5

# Сколько часов использовать сохраненный ответ из кэша.
CACHE_TTL_HOURS = 24

# True = игнорировать кэш и обновить всё заново.
FORCE_REFRESH = False

# В агрегированном списке конкурентов учитывать только товары,
# у которых была хотя бы одна ненулевая позиция за период.
ONLY_VISIBLE_COMPETITORS = True

MPSTATS_BASE = "https://mpstats.io"
SERP_URL_TEMPLATE = MPSTATS_BASE + "/api/oz/keywords/{query}/serp?fbs=false"


# ============================================================
# МОДЕЛИ
# ============================================================

@dataclass(frozen=True)
class QueryRelation:
    article: str
    own_sku: str
    query: str
    query_key: str
    frequency: int
    comparable: str
    thematic: str


# ============================================================
# УТИЛИТЫ
# ============================================================

def normalize_query(value: Any) -> str:
    text = "" if value is None else str(value)
    text = re.sub(r"\s+", " ", text.strip())
    return text.casefold()


def display_query(value: Any) -> str:
    text = "" if value is None else str(value)
    return re.sub(r"\s+", " ", text.strip())


def clean_id(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, bool):
        return str(value)
    if isinstance(value, int):
        return str(value)
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


def to_int(value: Any, default: int = 0) -> int:
    if value is None or value == "":
        return default
    try:
        return int(float(value))
    except (TypeError, ValueError):
        return default


def to_float(value: Any) -> float | None:
    if value is None or value == "":
        return None
    try:
        return float(value)
    except (TypeError, ValueError):
        return None


def safe_sum(values: Any) -> float:
    if not isinstance(values, (list, tuple)):
        return 0.0
    total = 0.0
    for v in values:
        try:
            total += float(v or 0)
        except (TypeError, ValueError):
            pass
    return total


def yes(value: Any) -> bool:
    return str(value or "").strip().casefold().startswith("да")


def utc_now_iso() -> str:
    return datetime.now(timezone.utc).isoformat()


# ============================================================
# ЧТЕНИЕ И ОТБОР ЗАПРОСОВ
# ============================================================

def load_input_rows(path: Path) -> list[dict[str, Any]]:
    if not path.exists():
        raise FileNotFoundError(
            f"Не найден входной файл:\n{path}\n\n"
            "Положите 'запросы 62 SKU.xlsx' в ту же папку, что и скрипт."
        )

    wb = load_workbook(path, read_only=True, data_only=True)
    ws = wb.active

    rows_iter = ws.iter_rows(values_only=True)
    headers = next(rows_iter)
    header_map = {str(v).strip(): idx for idx, v in enumerate(headers) if v is not None}

    required = ["Артикул", "SKU", "Запрос", "Частота из выгрузки"]
    missing = [h for h in required if h not in header_map]
    if missing:
        raise ValueError(f"Во входном файле не найдены колонки: {', '.join(missing)}")

    def get(row: tuple, name: str, default: Any = "") -> Any:
        idx = header_map.get(name)
        return row[idx] if idx is not None and idx < len(row) else default

    result: list[dict[str, Any]] = []
    for row in rows_iter:
        query = display_query(get(row, "Запрос"))
        sku = clean_id(get(row, "SKU"))
        if not query or not sku:
            continue

        result.append({
            "article": clean_id(get(row, "Артикул")),
            "sku": sku,
            "query": query,
            "frequency": to_int(get(row, "Частота из выгрузки")),
            "comparable": str(get(row, "Сопоставим: ≥7 дней в каждом", "") or ""),
            "thematic": str(get(row, "Тематический кандидат", "") or ""),
        })

    wb.close()
    return result


def select_query_relations(rows: list[dict[str, Any]]) -> list[QueryRelation]:
    """
    1. Дедуп по (SKU, нормализованный запрос).
    2. Фильтр тематичности / сопоставимости.
    3. TOP-N по частоте внутри каждого SKU.
    4. Если фильтр уничтожил все запросы SKU — fallback на исходный список.
    """
    all_by_sku: dict[str, dict[str, QueryRelation]] = defaultdict(dict)

    for row in rows:
        q_key = normalize_query(row["query"])
        if not q_key:
            continue

        rel = QueryRelation(
            article=row["article"],
            own_sku=row["sku"],
            query=row["query"],
            query_key=q_key,
            frequency=row["frequency"],
            comparable=row["comparable"],
            thematic=row["thematic"],
        )

        old = all_by_sku[rel.own_sku].get(q_key)
        if old is None or rel.frequency > old.frequency:
            all_by_sku[rel.own_sku][q_key] = rel

    selected: list[QueryRelation] = []

    for sku, qmap in all_by_sku.items():
        candidates = list(qmap.values())

        filtered = candidates
        if THEMATIC_ONLY:
            filtered = [r for r in filtered if yes(r.thematic)]
        if COMPARABLE_ONLY:
            filtered = [r for r in filtered if yes(r.comparable)]

        # Не оставляем SKU вообще без запросов.
        if not filtered:
            filtered = candidates

        filtered.sort(key=lambda r: (-r.frequency, r.query_key))

        if TOP_QUERIES_PER_SKU > 0:
            filtered = filtered[:TOP_QUERIES_PER_SKU]

        selected.extend(filtered)

    selected.sort(key=lambda r: (r.own_sku, -r.frequency, r.query_key))
    return selected


# ============================================================
# АВТОРИЗАЦИЯ ЧЕРЕЗ ОТДЕЛЬНЫЙ CHROME-ПРОФИЛЬ
# ============================================================

def get_browser_session_cookies() -> tuple[dict[str, str], str]:
    CHROME_PROFILE_DIR.mkdir(parents=True, exist_ok=True)

    options = Options()
    options.add_argument(f"--user-data-dir={CHROME_PROFILE_DIR}")
    options.add_argument("--start-maximized")

    print("\n[AUTH] Открываю Chrome с отдельным профилем MPSTATS...")
    driver = webdriver.Chrome(options=options)

    try:
        driver.get(MPSTATS_BASE + "/")

        print(
            "[AUTH] Если это первый запуск — войдите в MPSTATS вручную "
            "в открывшемся окне.\n"
            "[AUTH] Скрипт продолжит автоматически, когда появится cookie mp_auth."
        )

        deadline = time.time() + 300  # 5 минут
        mp_auth_found = False

        while time.time() < deadline:
            cookie = driver.get_cookie("mp_auth")
            if cookie and cookie.get("value"):
                mp_auth_found = True
                break
            time.sleep(1)

        if not mp_auth_found:
            raise RuntimeError(
                "За 5 минут не появилась cookie mp_auth. "
                "Проверьте, что вы вошли именно на https://mpstats.io/"
            )

        cookies = {c["name"]: c["value"] for c in driver.get_cookies()}
        user_agent = driver.execute_script("return navigator.userAgent;")

        print("[AUTH] Сессия MPSTATS найдена.")
        return cookies, user_agent

    finally:
        driver.quit()


# ============================================================
# HTTP + КЭШ
# ============================================================

_thread_local = threading.local()
_auth_cookies: dict[str, str] = {}
_user_agent = ""


def get_thread_session() -> requests.Session:
    session = getattr(_thread_local, "session", None)
    if session is None:
        session = requests.Session()
        session.cookies.update(_auth_cookies)
        session.headers.update({
            "Accept": "application/json, text/plain, */*",
            "Accept-Language": "ru,en-US;q=0.9,en;q=0.8",
            "Referer": MPSTATS_BASE + "/",
            "User-Agent": _user_agent,
        })
        _thread_local.session = session
    return session


def init_cache() -> sqlite3.Connection:
    conn = sqlite3.connect(CACHE_DB)
    conn.execute("""
        CREATE TABLE IF NOT EXISTS query_cache (
            query_key TEXT PRIMARY KEY,
            query_text TEXT NOT NULL,
            fetched_at TEXT NOT NULL,
            status INTEGER NOT NULL,
            response_json TEXT,
            error TEXT
        )
    """)
    conn.commit()
    return conn


def cache_get(conn: sqlite3.Connection, query_key: str) -> dict[str, Any] | None:
    if FORCE_REFRESH:
        return None

    row = conn.execute(
        "SELECT fetched_at, status, response_json FROM query_cache WHERE query_key = ?",
        (query_key,)
    ).fetchone()

    if not row:
        return None

    fetched_at, status, response_json = row
    if status != 200 or not response_json:
        return None

    try:
        fetched_dt = datetime.fromisoformat(fetched_at)
        age_hours = (datetime.now(timezone.utc) - fetched_dt).total_seconds() / 3600
        if age_hours > CACHE_TTL_HOURS:
            return None
        return json.loads(response_json)
    except Exception:
        return None


def cache_put(
    conn: sqlite3.Connection,
    query_key: str,
    query_text: str,
    status: int,
    data: dict[str, Any] | None,
    error: str | None = None,
) -> None:
    conn.execute("""
        INSERT INTO query_cache(query_key, query_text, fetched_at, status, response_json, error)
        VALUES (?, ?, ?, ?, ?, ?)
        ON CONFLICT(query_key) DO UPDATE SET
            query_text = excluded.query_text,
            fetched_at = excluded.fetched_at,
            status = excluded.status,
            response_json = excluded.response_json,
            error = excluded.error
    """, (
        query_key,
        query_text,
        utc_now_iso(),
        status,
        json.dumps(data, ensure_ascii=False) if data is not None else None,
        error,
    ))
    conn.commit()


def fetch_serp(query_text: str, query_key: str) -> tuple[str, str, int, dict[str, Any] | None, str | None]:
    """
    Возвращает:
    (query_key, query_text, http_status, json_or_none, error_or_none)
    """
    time.sleep(BASE_REQUEST_DELAY)

    encoded = quote(query_text, safe="")
    url = SERP_URL_TEMPLATE.format(query=encoded)

    session = get_thread_session()
    last_error: str | None = None
    last_status = 0

    for attempt in range(1, MAX_RETRIES + 1):
        try:
            response = session.get(url, timeout=40)
            last_status = response.status_code

            if response.status_code == 200:
                try:
                    data = response.json()
                except ValueError as exc:
                    return query_key, query_text, 200, None, f"Ответ не JSON: {exc}"
                return query_key, query_text, 200, data, None

            if response.status_code in (401, 403):
                return (
                    query_key, query_text, response.status_code, None,
                    "Сессия MPSTATS недействительна. Перезапустите скрипт и войдите заново."
                )

            if response.status_code == 429:
                retry_after = response.headers.get("Retry-After")
                try:
                    wait_s = max(float(retry_after), 1.0) if retry_after else 2 ** attempt
                except ValueError:
                    wait_s = 2 ** attempt

                last_error = f"429 Too Many Requests; ожидание {wait_s:.1f} сек."
                time.sleep(wait_s)
                continue

            if 500 <= response.status_code < 600:
                wait_s = 2 ** attempt
                last_error = f"HTTP {response.status_code}; повтор через {wait_s} сек."
                time.sleep(wait_s)
                continue

            text = response.text[:500].replace("\n", " ")
            return (
                query_key, query_text, response.status_code, None,
                f"HTTP {response.status_code}: {text}"
            )

        except requests.RequestException as exc:
            last_error = str(exc)
            time.sleep(2 ** attempt)

    return query_key, query_text, last_status, None, last_error or "Неизвестная ошибка"


# ============================================================
# ПАРСИНГ MPSTATS
# ============================================================

def parse_item(item: dict[str, Any], periods: list[str], response_rank: int) -> dict[str, Any]:
    positions = item.get("positions") or []

    nonzero_pairs: list[tuple[str | None, float]] = []
    for idx, raw_pos in enumerate(positions):
        pos = to_float(raw_pos)
        if pos is not None and pos > 0:
            date = periods[idx] if idx < len(periods) else None
            nonzero_pairs.append((date, pos))

    latest_day_position = None
    if positions:
        p = to_float(positions[-1])
        if p is not None and p > 0:
            latest_day_position = p

    last_seen_date = None
    last_seen_position = None
    best_position = None
    avg_position = None

    if nonzero_pairs:
        last_seen_date, last_seen_position = nonzero_pairs[-1]
        pos_values = [p for _, p in nonzero_pairs]
        best_position = min(pos_values)
        avg_position = mean(pos_values)

    advert = item.get("advert")
    ad_days = 0
    if isinstance(advert, dict):
        ad_days = sum(1 for v in advert.values() if to_float(v) and to_float(v) > 0)
    elif isinstance(advert, list):
        ad_days = sum(1 for v in advert if to_float(v) and to_float(v) > 0)

    return {
        "competitor_sku": clean_id(item.get("id")),
        "name": item.get("name") or "",
        "brand": item.get("brand") or "",
        "seller": item.get("seller") or "",
        "category": item.get("category") or item.get("niche") or "",
        "niche_id": clean_id(item.get("niche_id")),
        "final_price": to_float(item.get("final_price")),
        "ozon_card_price": to_float(item.get("ozon_card_price")),
        "rating": to_float(item.get("rating")),
        "comments": to_int(item.get("comments")),
        "orders_30d": safe_sum(item.get("graph")),
        "revenue_30d": safe_sum(item.get("revenue_graph")),
        "days_in_search": len(nonzero_pairs),
        "latest_day_position": latest_day_position,
        "last_seen_position": last_seen_position,
        "last_seen_date": last_seen_date,
        "best_position": best_position,
        "avg_position": avg_position,
        "ad_days": ad_days,
        "is_fbs": item.get("is_fbs"),
        "url": item.get("url") or "",
        "thumb": item.get("thumb") or "",
        "response_rank": response_rank,
    }


def parse_serp_response(data: dict[str, Any]) -> list[dict[str, Any]]:
    periods = data.get("periods") or []
    items = data.get("items") or []

    if MAX_ITEMS_PER_QUERY > 0:
        items = items[:MAX_ITEMS_PER_QUERY]

    parsed: list[dict[str, Any]] = []
    for idx, item in enumerate(items, start=1):
        if isinstance(item, dict):
            parsed.append(parse_item(item, periods, idx))

    return parsed


# ============================================================
# АГРЕГАЦИЯ КОНКУРЕНТОВ
# ============================================================

def freshness_factor(last_seen_date: str | None, all_periods_end: str | None) -> float:
    """
    Мягко понижает вес давно исчезнувших из выдачи товаров.
    Минимальный коэффициент 0.25.
    """
    if not last_seen_date or not all_periods_end:
        return 0.25

    try:
        last_seen = datetime.fromisoformat(last_seen_date).date()
        period_end = datetime.fromisoformat(all_periods_end).date()
        days = max((period_end - last_seen).days, 0)
        return max(0.25, 1.0 - min(days, 30) / 30.0)
    except Exception:
        return 0.25


def aggregate_competitors(
    selected: list[QueryRelation],
    query_items: dict[str, list[dict[str, Any]]],
    query_period_end: dict[str, str | None],
    own_skus: set[str],
) -> list[dict[str, Any]]:

    selected_by_sku: dict[str, list[QueryRelation]] = defaultdict(list)
    article_by_sku: dict[str, str] = {}

    for rel in selected:
        selected_by_sku[rel.own_sku].append(rel)
        article_by_sku.setdefault(rel.own_sku, rel.article)

    output: list[dict[str, Any]] = []

    for own_sku, relations in selected_by_sku.items():
        agg: dict[str, dict[str, Any]] = {}

        for rel in relations:
            items = query_items.get(rel.query_key, [])
            period_end = query_period_end.get(rel.query_key)

            for item in items:
                comp_sku = item["competitor_sku"]
                if not comp_sku or comp_sku in own_skus:
                    continue

                if ONLY_VISIBLE_COMPETITORS and item["days_in_search"] <= 0:
                    continue

                record = agg.get(comp_sku)
                if record is None:
                    record = {
                        "article": article_by_sku.get(own_sku, ""),
                        "own_sku": own_sku,
                        "competitor_sku": comp_sku,
                        "name": item["name"],
                        "brand": item["brand"],
                        "seller": item["seller"],
                        "category": item["category"],
                        "final_price": item["final_price"],
                        "ozon_card_price": item["ozon_card_price"],
                        "rating": item["rating"],
                        "comments": item["comments"],
                        "orders_30d": item["orders_30d"],
                        "revenue_30d": item["revenue_30d"],
                        "url": item["url"],
                        "query_hits": 0,
                        "sum_frequency": 0,
                        "score": 0.0,
                        "positions": [],
                        "best_position": None,
                        "last_seen_dates": [],
                        "queries": [],
                    }
                    agg[comp_sku] = record

                # Статические метрики товара не суммируем по запросам.
                # Берем максимум для продаж/выручки/отзывов на случай расхождений.
                record["comments"] = max(record["comments"], item["comments"])
                record["orders_30d"] = max(record["orders_30d"], item["orders_30d"])
                record["revenue_30d"] = max(record["revenue_30d"], item["revenue_30d"])

                record["query_hits"] += 1
                record["sum_frequency"] += rel.frequency
                record["queries"].append(rel.query)

                pos = item["last_seen_position"] or item["avg_position"]
                if pos:
                    record["positions"].append(float(pos))
                    if record["best_position"] is None:
                        record["best_position"] = float(item["best_position"] or pos)
                    else:
                        record["best_position"] = min(
                            record["best_position"],
                            float(item["best_position"] or pos)
                        )

                    freshness = freshness_factor(item["last_seen_date"], period_end)
                    # Эвристический балл:
                    # частотность запроса / sqrt(позиции) * свежесть присутствия.
                    record["score"] += (
                        rel.frequency / math.sqrt(max(float(pos), 1.0))
                    ) * freshness

                if item["last_seen_date"]:
                    record["last_seen_dates"].append(item["last_seen_date"])

        total_queries = len(relations)
        sku_rows: list[dict[str, Any]] = []

        for record in agg.values():
            positions = record.pop("positions")
            last_seen_dates = record.pop("last_seen_dates")
            queries = list(dict.fromkeys(record.pop("queries")))

            record["total_own_queries"] = total_queries
            record["query_share_pct"] = (
                round(record["query_hits"] / total_queries * 100, 1)
                if total_queries else 0
            )
            record["avg_position"] = round(mean(positions), 2) if positions else None
            record["best_position"] = (
                round(record["best_position"], 2)
                if record["best_position"] is not None else None
            )
            record["last_seen_date"] = max(last_seen_dates) if last_seen_dates else None
            record["score"] = round(record["score"], 2)
            record["queries_text"] = " | ".join(queries)
            sku_rows.append(record)

        sku_rows.sort(
            key=lambda r: (
                -r["score"],
                -r["query_hits"],
                -r["sum_frequency"],
                r["competitor_sku"],
            )
        )

        if TOP_COMPETITORS_PER_SKU > 0:
            sku_rows = sku_rows[:TOP_COMPETITORS_PER_SKU]

        for rank, row in enumerate(sku_rows, start=1):
            row["competitor_rank"] = rank
            output.append(row)

    output.sort(key=lambda r: (r["own_sku"], r["competitor_rank"]))
    return output


# ============================================================
# EXCEL OUTPUT
# ============================================================

HEADER_FILL = PatternFill("solid", fgColor="1F4E78")
HEADER_FONT = Font(color="FFFFFF", bold=True)


def style_sheet(ws, widths: dict[int, int] | None = None) -> None:
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions

    for cell in ws[1]:
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

    if widths:
        for col_idx, width in widths.items():
            ws.column_dimensions[get_column_letter(col_idx)].width = width


def write_output(
    selected: list[QueryRelation],
    query_items: dict[str, list[dict[str, Any]]],
    query_period_end: dict[str, str | None],
    aggregated: list[dict[str, Any]],
    errors: list[tuple[str, int, str]],
    cached_count: int,
    fetched_count: int,
) -> None:

    wb = Workbook()

    # 1. Сводка
    ws = wb.active
    ws.title = "Сводка"
    summary_rows = [
        ["Параметр", "Значение"],
        ["Дата формирования", datetime.now().strftime("%Y-%m-%d %H:%M:%S")],
        ["Входной файл", INPUT_XLSX.name],
        ["TOP запросов на SKU", TOP_QUERIES_PER_SKU or "Все"],
        ["Только тематические", THEMATIC_ONLY],
        ["Только сопоставимые", COMPARABLE_ONLY],
        ["Параллельных запросов", MAX_WORKERS],
        ["Выбрано связей SKU ↔ запрос", len(selected)],
        ["Уникальных запросов", len({r.query_key for r in selected})],
        ["Из кэша", cached_count],
        ["Получено с MPSTATS", fetched_count],
        ["Ошибок", len(errors)],
        ["Агрегированных конкурентов", len(aggregated)],
        ["Кэш", CACHE_DB.name],
    ]
    for row in summary_rows:
        ws.append(row)
    ws["A1"].fill = HEADER_FILL
    ws["A1"].font = HEADER_FONT
    ws["B1"].fill = HEADER_FILL
    ws["B1"].font = HEADER_FONT
    ws.column_dimensions["A"].width = 32
    ws.column_dimensions["B"].width = 60

    # 2. Выбранные запросы
    ws = wb.create_sheet("Запросы")
    ws.append([
        "Артикул", "Наш SKU", "Запрос", "Частота",
        "Сопоставим ≥7 дней", "Тематический кандидат"
    ])
    for rel in selected:
        ws.append([
            rel.article, rel.own_sku, rel.query, rel.frequency,
            rel.comparable, rel.thematic
        ])
    style_sheet(ws, {1: 16, 2: 16, 3: 42, 4: 12, 5: 20, 6: 34})

    # 3. Сырая выдача — один запрос × один конкурент
    ws = wb.create_sheet("Выдача")
    raw_headers = [
        "Запрос", "SKU конкурента", "Название", "Бренд", "Продавец",
        "Категория", "Цена", "Цена Ozon", "Рейтинг", "Отзывы",
        "Заказы 30д", "Выручка 30д", "Дней в выдаче",
        "Позиция в последний день", "Последняя видимая позиция",
        "Дата последнего появления", "Лучшая позиция", "Средняя позиция",
        "Дней в рекламе", "FBS", "URL", "№ в ответе"
    ]
    ws.append(raw_headers)

    query_display = {}
    for rel in selected:
        query_display.setdefault(rel.query_key, rel.query)

    for q_key in sorted(query_items, key=lambda k: query_display.get(k, k)):
        q_text = query_display.get(q_key, q_key)
        for item in query_items[q_key]:
            ws.append([
                q_text,
                item["competitor_sku"],
                item["name"],
                item["brand"],
                item["seller"],
                item["category"],
                item["final_price"],
                item["ozon_card_price"],
                item["rating"],
                item["comments"],
                item["orders_30d"],
                item["revenue_30d"],
                item["days_in_search"],
                item["latest_day_position"],
                item["last_seen_position"],
                item["last_seen_date"],
                item["best_position"],
                round(item["avg_position"], 2) if item["avg_position"] is not None else None,
                item["ad_days"],
                item["is_fbs"],
                item["url"],
                item["response_rank"],
            ])

    style_sheet(ws, {
        1: 38, 2: 16, 3: 52, 4: 22, 5: 24, 6: 48,
        7: 12, 8: 12, 9: 10, 10: 10, 11: 12, 12: 14,
        13: 13, 14: 15, 15: 18, 16: 18, 17: 13, 18: 14,
        19: 13, 20: 8, 21: 45, 22: 11
    })

    # 4. Агрегированные конкуренты
    ws = wb.create_sheet("Конкуренты")
    agg_headers = [
        "Артикул", "Наш SKU", "Место конкурента", "SKU конкурента",
        "Название", "Бренд", "Продавец", "Категория",
        "Цена", "Цена Ozon", "Рейтинг", "Отзывы",
        "Заказы 30д", "Выручка 30д",
        "Пересечение запросов", "Запросов у нашего SKU", "Доля запросов, %",
        "Сумма частотности", "Лучшая позиция", "Средняя позиция",
        "Последнее появление", "Конкурентный балл", "Запросы", "URL"
    ]
    ws.append(agg_headers)

    for r in aggregated:
        ws.append([
            r["article"],
            r["own_sku"],
            r["competitor_rank"],
            r["competitor_sku"],
            r["name"],
            r["brand"],
            r["seller"],
            r["category"],
            r["final_price"],
            r["ozon_card_price"],
            r["rating"],
            r["comments"],
            r["orders_30d"],
            r["revenue_30d"],
            r["query_hits"],
            r["total_own_queries"],
            r["query_share_pct"],
            r["sum_frequency"],
            r["best_position"],
            r["avg_position"],
            r["last_seen_date"],
            r["score"],
            r["queries_text"],
            r["url"],
        ])

    style_sheet(ws, {
        1: 15, 2: 16, 3: 12, 4: 16, 5: 52, 6: 22, 7: 24, 8: 42,
        9: 12, 10: 12, 11: 10, 12: 10, 13: 12, 14: 14,
        15: 16, 16: 18, 17: 14, 18: 16, 19: 14, 20: 15,
        21: 18, 22: 18, 23: 70, 24: 45
    })

    # 5. Ошибки
    ws = wb.create_sheet("Ошибки")
    ws.append(["Запрос", "HTTP", "Ошибка"])
    for query, status, error in errors:
        ws.append([query, status, error])
    style_sheet(ws, {1: 45, 2: 10, 3: 100})

    # Форматы чисел
    for sheet_name in ("Выдача", "Конкуренты"):
        sh = wb[sheet_name]
        for row in sh.iter_rows(min_row=2):
            for cell in row:
                if isinstance(cell.value, float):
                    cell.number_format = '#,##0.00'

    wb.save(OUTPUT_XLSX)


# ============================================================
# MAIN
# ============================================================

def main() -> None:
    print("=" * 70)
    print("MPSTATS — автоматический сбор конкурентов Ozon")
    print("=" * 70)

    print(f"\n[1/6] Читаю {INPUT_XLSX.name}...")
    input_rows = load_input_rows(INPUT_XLSX)
    own_skus = {row["sku"] for row in input_rows}

    selected = select_query_relations(input_rows)
    unique_queries: dict[str, str] = {}
    for rel in selected:
        unique_queries.setdefault(rel.query_key, rel.query)

    print(f"Строк во входном файле: {len(input_rows):,}".replace(",", " "))
    print(f"Наших SKU: {len(own_skus)}")
    print(f"Выбрано связей SKU ↔ запрос: {len(selected)}")
    print(f"Уникальных запросов к MPSTATS: {len(unique_queries)}")

    print("\n[2/6] Получаю авторизованную сессию...")
    global _auth_cookies, _user_agent
    _auth_cookies, _user_agent = get_browser_session_cookies()

    if "mp_auth" not in _auth_cookies:
        raise RuntimeError("Cookie mp_auth не найдена.")

    print("\n[3/6] Проверяю кэш...")
    cache_conn = init_cache()

    responses: dict[str, dict[str, Any]] = {}
    to_fetch: list[tuple[str, str]] = []
    cached_count = 0

    for q_key, q_text in unique_queries.items():
        cached = cache_get(cache_conn, q_key)
        if cached is not None:
            responses[q_key] = cached
            cached_count += 1
        else:
            to_fetch.append((q_key, q_text))

    print(f"Из кэша: {cached_count}")
    print(f"Нужно запросить: {len(to_fetch)}")

    errors: list[tuple[str, int, str]] = []
    fetched_count = 0

    if to_fetch:
        print(f"\n[4/6] Запрашиваю MPSTATS ({MAX_WORKERS} потока)...")

        with ThreadPoolExecutor(max_workers=MAX_WORKERS) as executor:
            futures = {
                executor.submit(fetch_serp, q_text, q_key): (q_key, q_text)
                for q_key, q_text in to_fetch
            }

            total = len(futures)
            completed = 0

            for future in as_completed(futures):
                q_key, q_text = futures[future]
                completed += 1

                try:
                    ret_key, ret_text, status, data, error = future.result()
                except Exception as exc:
                    status = 0
                    data = None
                    error = repr(exc)
                    ret_key, ret_text = q_key, q_text

                if data is not None and status == 200:
                    responses[ret_key] = data
                    cache_put(cache_conn, ret_key, ret_text, status, data, None)
                    fetched_count += 1
                    item_count = len(data.get("items") or [])
                    print(f"[{completed}/{total}] OK  {ret_text!r} → {item_count} товаров")
                else:
                    err = error or "Неизвестная ошибка"
                    cache_put(cache_conn, ret_key, ret_text, status, None, err)
                    errors.append((ret_text, status, err))
                    print(f"[{completed}/{total}] ERR {ret_text!r} → HTTP {status}: {err}")

                    if status in (401, 403):
                        print(
                            "\nСессия MPSTATS истекла. Уже полученные ответы сохранены в кэш.\n"
                            "Запустите скрипт снова — он продолжит с оставшихся запросов."
                        )
                        break
    else:
        print("\n[4/6] Все запросы уже есть в свежем кэше.")

    cache_conn.close()

    print("\n[5/6] Разбираю выдачу и считаю конкурентов...")
    query_items: dict[str, list[dict[str, Any]]] = {}
    query_period_end: dict[str, str | None] = {}

    for q_key, data in responses.items():
        periods = data.get("periods") or []
        query_period_end[q_key] = periods[-1] if periods else None
        query_items[q_key] = parse_serp_response(data)

    aggregated = aggregate_competitors(
        selected=selected,
        query_items=query_items,
        query_period_end=query_period_end,
        own_skus=own_skus,
    )

    print(f"Сформировано строк конкурентов: {len(aggregated)}")

    print("\n[6/6] Формирую Excel...")
    write_output(
        selected=selected,
        query_items=query_items,
        query_period_end=query_period_end,
        aggregated=aggregated,
        errors=errors,
        cached_count=cached_count,
        fetched_count=fetched_count,
    )

    print("\nГОТОВО")
    print(f"Результат: {OUTPUT_XLSX}")
    print(f"Кэш:      {CACHE_DB}")
    print(
        "\nПовторный запуск использует кэш и не будет заново дергать "
        "свежие запросы MPSTATS."
    )


if __name__ == "__main__":
    try:
        main()
    except KeyboardInterrupt:
        print("\n\nОстановлено пользователем. Уже полученные ответы остались в кэше.")
    except Exception as exc:
        print("\nОШИБКА:")
        print(exc)
        print(
            "\nЕсли ошибка связана с авторизацией, просто запустите скрипт снова "
            "и войдите в MPSTATS в открывшемся Chrome."
        )
        raise
