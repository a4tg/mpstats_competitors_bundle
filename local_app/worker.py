"""Run existing collector with the selected SKU set, isolated output and cache."""
import json
import sys
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
import mpstats_competitors as collector


def main():
    folder = Path(sys.argv[1]).resolve()
    config = json.loads((folder / "files.json").read_text(encoding="utf-8"))
    data = json.loads((folder / "data.json").read_text(encoding="utf-8"))
    selected = {p["sku"] for p in data["products"] if p["sku"]}
    reader = collector.load_input_rows
    source = folder / config["queries"]
    all_rows = reader(source)
    rows = [row for row in all_rows if row["sku"] in selected]
    missing = selected - {row["sku"] for row in rows}
    if missing:
        raise ValueError("В таблице запросов отсутствуют SKU: " + ", ".join(sorted(missing)))
    if not rows:
        raise ValueError("Нет запросов для выбранных товаров")
    collector.load_input_rows = lambda _: rows
    # Preserve every own SKU present in the query source during exclusion.
    original_aggregate = collector.aggregate_competitors
    all_own = selected | {row["sku"] for row in all_rows}
    def aggregate(**kwargs):
        kwargs["own_skus"] = all_own
        return original_aggregate(**kwargs)
    collector.aggregate_competitors = aggregate
    collector.INPUT_XLSX = source
    collector.OUTPUT_XLSX = folder / "competitors.xlsx"
    collector.CACHE_DB = folder / "cache.sqlite3"
    collector.main()
    from openpyxl import load_workbook
    workbook = load_workbook(collector.OUTPUT_XLSX, read_only=True, data_only=True)
    try:
        summary = dict(workbook["Сводка"].iter_rows(min_row=2, max_col=2, values_only=True))
        total = summary.get("Уникальных запросов", 0)
        obtained = summary.get("Из кэша", 0) + summary.get("Получено с MPSTATS", 0)
        if summary.get("Ошибок", 0) or obtained != total:
            raise RuntimeError(f"Неполный сбор: получено {obtained}/{total}. Повторите сбор; ответы сохранены в кэше.")
    finally:
        workbook.close()


if __name__ == "__main__":
    main()
