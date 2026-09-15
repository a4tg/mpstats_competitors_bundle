import hashlib
import json
import re
import shutil
import time
import uuid
from datetime import date, datetime
from pathlib import Path
from urllib.parse import urlencode
from zipfile import ZipFile

from openpyxl import load_workbook
from selenium import webdriver
from selenium.webdriver import ActionChains
from selenium.webdriver.common.by import By
from selenium.webdriver.support.ui import WebDriverWait


BASE = Path(
    r"C:\Users\artem\Desktop\Прямой Контракт\Поисковые запросы"
)
DATE_FROM = "17.08.2026"
DATE_TO = "13.09.2026"
PERIOD = f"{DATE_FROM}_{DATE_TO}"
STATE_FILE = BASE / "progress.json"

# Селекторы предполагают AG Grid, как у таблицы на скриншоте.
# Если структура MPStats отличается, скрипт предложит ручной экспорт.
GRID_SELECTOR = ".ag-root"
MENU_SELECTOR = ".ag-menu-option, [role='menuitem']"


def digest(path):
    return hashlib.sha256(path.read_bytes()).hexdigest()


def valid_export(path):
    """Проверяет XLSX, заголовок запросов и границы периода."""
    with ZipFile(path) as archive:
        if archive.testzip() is not None:
            raise RuntimeError("Повреждённый XLSX")

    book = load_workbook(path, read_only=True, data_only=True)
    header_text = []
    try:
        for sheet in book:
            sheet.reset_dimensions()
            for row in sheet.iter_rows(
                min_row=1, max_row=15, values_only=True
            ):
                for value in row:
                    if value is None:
                        continue
                    if isinstance(value, (datetime, date)):
                        header_text.append(value.strftime("%d.%m.%Y"))
                    else:
                        header_text.append(str(value))
    finally:
        book.close()

    text = " ".join(header_text).lower()
    if "запрос" not in text:
        raise RuntimeError("В XLSX не найден заголовок запросов")

    for boundary in (DATE_FROM, DATE_TO):
        parsed = datetime.strptime(boundary, "%d.%m.%Y")
        variants = (
            boundary,
            parsed.strftime("%Y-%m-%d"),
            parsed.strftime("%d/%m/%Y"),
        )
        if not any(value in text for value in variants):
            raise RuntimeError(
                f"В заголовках XLSX не найдена дата {boundary}. "
                "Проверь период и полноту экспорта. "
                "Скачанный файл оставлен во временной папке."
            )


def find_query_cell(driver):
    # Выбираем именно таблицу запросов, не соседние таблицы.
    for grid in driver.find_elements(By.CSS_SELECTOR, GRID_SELECTOR):
        if not grid.is_displayed():
            continue

        headers = grid.find_elements(
            By.CSS_SELECTOR, ".ag-header-cell"
        )
        query_header = next(
            (
                header for header in headers
                if header.text.strip().lower() == "запрос"
            ),
            None,
        )
        if query_header is None:
            continue

        column = query_header.get_attribute("col-id")
        if not column:
            continue

        for cell in grid.find_elements(By.CSS_SELECTOR, ".ag-cell"):
            if (
                cell.get_attribute("col-id") == column
                and cell.is_displayed()
                and cell.text.strip()
            ):
                return cell
    return False


def find_menu_item(driver, predicate):
    for item in driver.find_elements(By.CSS_SELECTOR, MENU_SELECTOR):
        if item.is_displayed() and predicate(item.text.strip().lower()):
            return item
    return False


def export_via_menu(driver):
    wait = WebDriverWait(driver, 35)
    cell = wait.until(find_query_cell)
    driver.execute_script(
        "arguments[0].scrollIntoView({block:'center'});", cell
    )
    ActionChains(driver).context_click(cell).perform()

    download = wait.until(
        lambda d: find_menu_item(
            d, lambda text: text in ("скачать", "экспорт", "export")
        )
    )

    # В AG Grid подменю обычно открывается наведением.
    ActionChains(driver).move_to_element(download).pause(1).perform()

    def excel_item(d):
        return find_menu_item(
            d, lambda text: "xlsx" in text or "excel" in text
        )

    try:
        excel = WebDriverWait(driver, 5).until(excel_item)
    except Exception:
        download.click()
        excel = WebDriverWait(driver, 10).until(excel_item)

    excel.click()


def wait_download(folder, timeout=120):
    deadline = time.monotonic() + timeout
    previous = None
    stable = 0

    while time.monotonic() < deadline:
        files = list(folder.iterdir())
        partial = any(
            p.suffix.lower() in (".crdownload", ".tmp", ".part")
            for p in files
        )
        candidates = [
            p for p in files if p.suffix.lower() == ".xlsx"
        ]

        if len(candidates) > 1:
            raise RuntimeError(
                f"В папке несколько XLSX, нельзя выбрать однозначно: {folder}"
            )

        if len(candidates) == 1 and not partial:
            file = candidates[0]
            signature = (file.name, file.stat().st_size)
            stable = stable + 1 if signature == previous else 0
            previous = signature

            if stable >= 2 and signature[1] > 0:
                return file
        else:
            stable = 0

        time.sleep(1)

    raise TimeoutError(f"Скачивание не завершилось: {folder}")


def main():
    BASE.mkdir(parents=True, exist_ok=True)
    sku_file = BASE / "sku.txt"

    if not sku_file.exists():
        raise RuntimeError(f"Создай файл со SKU: {sku_file}")

    skus = list(dict.fromkeys(
        line.strip()
        for line in sku_file.read_text(encoding="utf-8-sig").splitlines()
        if line.strip()
    ))
    if not skus or any(not re.fullmatch(r"\d+", sku) for sku in skus):
        raise RuntimeError(
            "В sku.txt должны быть только числовые SKU, по одному на строку"
        )

    state = (
        json.loads(STATE_FILE.read_text(encoding="utf-8"))
        if STATE_FILE.exists() else {}
    )

    options = webdriver.ChromeOptions()
    # Отдельный профиль. В свой обычный Chrome скрипт не вмешивается.
    options.add_argument(
        f"--user-data-dir={BASE / 'chrome_mpstats_profile'}"
    )
    options.add_argument("--start-maximized")
    options.add_experimental_option("prefs", {
        "download.default_directory": str(BASE),
        "download.prompt_for_download": False,
        "download.directory_upgrade": True,
    })

    driver = webdriver.Chrome(options=options)
    driver.set_page_load_timeout(60)

    try:
        driver.get("https://mpstats.io/")
        input(
            "\nВойди в MPStats в открывшемся Chrome. "
            "После входа нажми Enter здесь..."
        )

        for index, sku in enumerate(skus, 1):
            target = BASE / f"{sku}.xlsx"
            key = f"{sku}:{PERIOD}"
            record = state.get(key)

            if (
                target.exists()
                and record
                and record.get("sha256") == digest(target)
            ):
                print(f"[{index}/{len(skus)} {sku}: уже сохранён")
                continue

            if target.exists():
                raise RuntimeError(
                    f"{target.name} уже существует, но не подтверждён "
                    "в progress.json за этот период. "
                    "Перемести его в другую папку перед повторным запуском."
                )

            # Отдельная папка исключает путаницу с предыдущими загрузками.
            temporary = BASE / "_downloads" / (
                f"{sku}_{uuid.uuid4().hex[:8]}"
            )
            temporary.mkdir(parents=True)

            driver.execute_cdp_cmd("Browser.setDownloadBehavior", {
                "behavior": "allow",
                "downloadPath": str(temporary),
                "eventsEnabled": True,
            })

            query = urlencode({
                "d1": DATE_FROM,
                "d2": DATE_TO,
                "tab": "visibility",
            })
            url = f"https://mpstats.io/ozon/item/{sku}?{query}"
            print(f"[{index}/{len(skus)}] Обрабатываю {sku}")
            driver.get(url)

            try:
                export_via_menu(driver)
            except Exception as error:
                screenshot = BASE / f"error_{sku}.png"
                driver.save_screenshot(str(screenshot))
                print(
                    f"\nНе удалось нажать экспорт: {type(error).__name__}."
                    f"\nСкриншот: {screenshot}"
                    "\nПроверь, что открыт нужный SKU и нужный период."
                    "\nМожно вручную выполнить ПКМ → Скачать → Excel."
                )
                answer = input(
                    "После ручного экспорта нажми Enter; "
                    "для остановки введи q: "
                )
                if answer.strip().lower() == "q":
                    raise RuntimeError(f"Остановлено на SKU {sku}")

            downloaded = wait_download(temporary)
            valid_export(downloaded)
            shutil.move(str(downloaded), str(target))

            state[key] = {
                "file": target.name,
                "sha256": digest(target),
                "url": url,
                "saved_at": datetime.now().isoformat(timespec="seconds"),
            }
            state_tmp = STATE_FILE.with_suffix(".tmp")
            state_tmp.write_text(
                json.dumps(state, ensure_ascii=False, indent=2),
                encoding="utf-8",
            )
            state_tmp.replace(STATE_FILE)

            print(f"Сохранён: {target.name}")
            time.sleep(3)

        print(f"\nГотово. Обработан весь список: {len(skus)} SKU.")

    finally:
        driver.quit()


if __name__ == "__main__":
    try:
        main()
    except Exception as error:
        print(f"\nОСТАНОВКА: {error}")
        print("Успешные выгрузки сохранены. Можно запустить снова.")
    input("\nНажми Enter для закрытия...")