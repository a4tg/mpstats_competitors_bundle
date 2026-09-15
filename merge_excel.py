from pathlib import Path
from copy import copy
from openpyxl import Workbook, load_workbook

FOLDER = Path(
    r"C:\Users\artem\Desktop\Прямой Контракт\Поисковые запросы\requests"
)
OUTPUT = FOLDER / "Все_запросы.xlsx"

# Берём только файлы, названные числовыми SKU.
files = sorted(
    p for p in FOLDER.glob("*.xlsx")
    if p.stem.isdigit()
)

if not files:
    raise SystemExit("Не найдены XLSX-файлы с числовыми SKU в названиях.")

# Проверяем имена заранее, чтобы не объединить разные SKU под одним именем.
names = [p.stem for p in files]
if any(len(name) > 31 for name in names):
    raise SystemExit("Название листа Excel не может превышать 31 символ.")

if OUTPUT.exists():
    raise SystemExit(
        f"Файл {OUTPUT.name} уже существует. Переименуй его перед запуском."
    )

result = Workbook()
result.remove(result.active)

for path in files:
    source = load_workbook(path)
    try:
        src = source.worksheets[0]
        dst = result.create_sheet(title=path.stem)

        for row in src.iter_rows():
            for cell in row:
                target = dst.cell(
                    row=cell.row,
                    column=cell.column,
                    value=cell.value,
                )
                if cell.has_style:
                    target._style = copy(cell._style)
                if cell.hyperlink:
                    target.hyperlink = copy(cell.hyperlink)
                if cell.comment:
                    target.comment = copy(cell.comment)

        for key, dimension in src.column_dimensions.items():
            dst.column_dimensions[key] = copy(dimension)

        for key, dimension in src.row_dimensions.items():
            dst.row_dimensions[key] = copy(dimension)

        for merged_range in src.merged_cells.ranges:
            dst.merge_cells(str(merged_range))

        dst.freeze_panes = src.freeze_panes
        dst.auto_filter.ref = src.auto_filter.ref

        print(f"Добавлен лист: {path.stem}")
    finally:
        source.close()

result.save(OUTPUT)
print(f"\nГотово: {OUTPUT}")
print(f"Листов: {len(files)}")